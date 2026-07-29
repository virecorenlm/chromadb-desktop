import PyPDF2
import docx
import os
import magic
from typing import List, Dict, Any
import logging

class FileParser:
    def __init__(self):
        self.logger = logging.getLogger(__name__)
        self.mime = magic.Magic(mime=True)
    
    def detect_file_type(self, file_path: str) -> str:
        """Detect file type using python-magic"""
        try:
            return self.mime.from_file(file_path)
        except Exception as e:
            self.logger.error(f"Error detecting file type: {e}")
            return "unknown"
    
    def parse_pdf(self, file_path: str) -> List[str]:
        """Extract text from PDF file"""
        try:
            text_chunks = []
            with open(file_path, 'rb') as file:
                reader = PyPDF2.PdfReader(file)
                for page_num in range(len(reader.pages)):
                    page = reader.pages[page_num]
                    text = page.extract_text()
                    if text.strip():
                        text_chunks.append(text)
            return text_chunks
        except Exception as e:
            self.logger.error(f"Error parsing PDF {file_path}: {e}")
            return []
    
    def parse_docx(self, file_path: str) -> List[str]:
        """Extract text from DOCX file"""
        try:
            doc = docx.Document(file_path)
            text_chunks = []
            current_chunk = ""
            
            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    current_chunk += paragraph.text + "\n"
                    # Simple chunking - you can make this smarter
                    if len(current_chunk) > 1000:
                        text_chunks.append(current_chunk)
                        current_chunk = ""
            
            if current_chunk:
                text_chunks.append(current_chunk)
                
            return text_chunks if text_chunks else [" ".join(p.text for p in doc.paragraphs if p.text)]
        except Exception as e:
            self.logger.error(f"Error parsing DOCX {file_path}: {e}")
            return []
    
    def parse_txt(self, file_path: str) -> List[str]:
        """Extract text from plain text file"""
        try:
            with open(file_path, 'r', encoding='utf-8', errors='ignore') as file:
                content = file.read()
                # Simple chunking - split by large paragraphs
                chunks = [chunk for chunk in content.split('\n\n') if chunk.strip()]
                return chunks if chunks else [content]
        except Exception as e:
            self.logger.error(f"Error parsing TXT {file_path}: {e}")
            return []
    
    def parse_file(self, file_path: str) -> Dict[str, Any]:
        """Main method to parse any supported file type"""
        if not os.path.exists(file_path):
            self.logger.error(f"File not found: {file_path}")
            return {"success": False, "chunks": [], "error": "File not found"}
        
        file_type = self.detect_file_type(file_path)
        filename = os.path.basename(file_path)
        
        self.logger.info(f"Parsing file: {filename}, Type: {file_type}")
        
        chunks = []
        
        if file_type == "application/pdf":
            chunks = self.parse_pdf(file_path)
        elif file_type in ["application/vnd.openxmlformats-officedocument.wordprocessingml.document", 
                          "application/msword"]:
            chunks = self.parse_docx(file_path)
        elif file_type.startswith("text/"):
            chunks = self.parse_txt(file_path)
        else:
            self.logger.warning(f"Unsupported file type: {file_type}")
            return {"success": False, "chunks": [], "error": f"Unsupported file type: {file_type}"}
        
        return {
            "success": True,
            "chunks": chunks,
            "filename": filename,
            "file_type": file_type,
            "chunk_count": len(chunks)
        }
